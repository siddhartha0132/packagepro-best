"""Builds docs/pitch/PackagePro-PS04.pptx — the 9-slide PS-04 pitch deck (editable shapes, speaker notes on every slide).

Usage:  python3 -m venv .venv && .venv/bin/pip install python-pptx && .venv/bin/python docs/pitch/build_deck.py
Every figure on the slides comes from the running app or its tests; business-model rates are left as [__] for the team.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).with_name("PackagePro-PS04.pptx")

NAVY, LIGHT, PAPER = "0B1F3A", "F4F7FB", "FBFCFE"
BLUE, ORANGE = "0B6BCB", "F59E5B"
INK, SOFT, MUTED_ON_DARK = "0B1F3A", "3F4C63", "BFD0E6"
CARD_DARK, LINE_DARK, LINE = "13294A", "22406B", "DCE3EE"
GREEN, GREEN_BG, GREEN_LINE = "0E8A5F", "E9F7F0", "B7EBD3"
RED, RED_BG, RED_LINE = "B42318", "FFF4F1", "F3C1B8"
HEAD, BODY = "Calibri", "Calibri"
W, H, M = 13.333, 7.5, 0.6
TOTAL = 9


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]


def slide(bg: str):
    s = prs.slides.add_slide(BLANK)
    fill = s.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(bg)
    return s


def text(s, x, y, w, h, lines, size=16, color=SOFT, bold=False, font=BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15):
    """lines: str | list of str | list of list[(text, {bold, color})] runs."""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        runs = line if isinstance(line, list) else [(line, {})]
        for chunk, style in runs:
            r = p.add_run()
            r.text = chunk
            r.font.name = style.get("font", font)
            r.font.size = Pt(style.get("size", size))
            r.font.bold = style.get("bold", bold)
            r.font.color.rgb = rgb(style.get("color", color))
    return box


def card(s, x, y, w, h, fill=PAPER, line=LINE, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    shp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    return shp


def pill(s, x, y, w, label, fill, color, size=11):
    card(s, x, y, w, 0.32, fill=fill, line=None, radius=0.5)
    text(s, x, y, w, 0.32, label, size=size, color=color, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def arrow(s, x, y, w=0.45, h=0.32, color=BLUE, shape=MSO_SHAPE.RIGHT_ARROW):
    a = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = rgb(color)
    a.line.fill.background()
    a.shadow.inherit = False
    return a


def header(s, eyebrow, title, dark=False):
    text(s, M, 0.5, W - 2 * M, 0.3, eyebrow.upper(), size=12, bold=True, color=ORANGE if dark else BLUE)
    text(s, M, 0.82, W - 2 * M, 0.7, title, size=32, bold=True, font=HEAD, color=LIGHT if dark else INK)


def footer(s, n, dark=False):
    c = MUTED_ON_DARK if dark else "6A7689"
    text(s, M, H - 0.45, 8, 0.25, "PackagePro · PS-04 Dynamic Tour Packages · Team RNG Gods", size=10, color=c)
    text(s, W - M - 1.5, H - 0.45, 1.5, 0.25, f"{n} / {TOTAL}", size=10, color=c, align=PP_ALIGN.RIGHT)


def notes(s, body):
    s.notes_slide.notes_text_frame.text = body


# ---------------------------------------------------------------- 1 · Problem statement
s = slide(NAVY)
text(s, M, 0.55, 11, 0.3, "PS-04 · KOGNIVERA HACKATHON 2026 · TEAM RNG GODS · BMS COLLEGE OF ENGINEERING", size=12, bold=True, color=ORANGE)
text(s, M, 0.95, 11, 1.2, "PackagePro", size=60, bold=True, font=HEAD, color=LIGHT)
text(s, M, 2.0, 11, 0.5, "Dynamic tour packages, priced live", size=24, color=MUTED_ON_DARK)
text(s, M, 2.8, W - 2 * M, 1.35, [[("The problem  ", {"bold": True, "color": ORANGE}), ("Holiday packages are sold as fixed bundles. Travellers can't reshape them, can't see what each part costs, and a local guide is picked from a static list — with no check that the guide is actually free on the trip's dates.", {})]], size=18, color=LIGHT, spacing=1.25)
problems = [
    ("Rigid bundles", "Changing the hotel, an activity or a transfer means calling an agent and waiting for a new quote."),
    ("Opaque pricing", "One lump sum. Nobody can see what each part costs or what a change would do to the total."),
    ("Guides picked blind", "A guide can be sold for a day they are already booked, and language or specialisation is ignored."),
]
cw = (W - 2 * M - 2 * 0.3) / 3
for i, (title, body) in enumerate(problems):
    x = M + i * (cw + 0.3)
    card(s, x, 4.4, cw, 2.2, fill=CARD_DARK, line=LINE_DARK)
    text(s, x + 0.3, 4.6, cw - 0.6, 0.4, f"0{i + 1}", size=14, bold=True, color=ORANGE)
    text(s, x + 0.3, 4.95, cw - 0.6, 0.45, title, size=20, bold=True, font=HEAD, color=LIGHT)
    text(s, x + 0.3, 5.45, cw - 0.6, 1.1, body, size=14, color=MUTED_ON_DARK, spacing=1.2)
footer(s, 1, dark=True)
notes(s, "PS-04 asks for dynamic tour packages. Today a package is a fixed bundle: you cannot swap parts, you cannot see the price of each part, and the local guide is a static dropdown with no availability check. Our job: make a package you can reshape, priced live, with guides that are checked day by day.")

# ---------------------------------------------------------------- 2 · Idea / proposed solution
s = slide(LIGHT)
header(s, "Idea · proposed solution", "A package you can reshape — every change repriced live")
text(s, M, 1.6, W - 2 * M, 0.6, "Start from a curated package, reshape it piece by piece, add a local guide that is checked against real availability, and book it — on the web or in Telegram, in your language.", size=16, color=SOFT)
pillars = [
    ("Curated start", "45 real PS-04 packages, browsed by theme and mood. Picks change with what you want: beaches, heritage, mountains."),
    ("Reshape live", "Swap the hotel, activity or transfer, add extras, change the days. The total is rebuilt from the whole plan on every change."),
    ("Guides checked day by day", "Pick by language, specialisation and price. A clash is refused with the date named; a matching substitute is offered and repriced."),
    ("Your language, your channel", "English, हिन्दी, தமிழ், తెలుగు. Web app, Telegram bot and a printable quotation, all on one engine."),
]
cw = (W - 2 * M - 3 * 0.25) / 4
for i, (title, body) in enumerate(pillars):
    x = M + i * (cw + 0.25)
    card(s, x, 2.45, cw, 2.9, radius=0.03)
    card(s, x, 2.45, cw, 0.09, fill=BLUE, line=None, radius=0, shape=MSO_SHAPE.RECTANGLE)
    text(s, x + 0.25, 2.75, cw - 0.5, 0.7, title, size=18, bold=True, font=HEAD, color=INK)
    text(s, x + 0.25, 3.5, cw - 0.5, 1.8, body, size=13, color=SOFT, spacing=1.2)
card(s, M, 5.6, W - 2 * M, 1.05, fill="EAF2FD", line="C9DDF7")
text(s, M + 0.3, 5.72, W - 2 * M - 0.6, 0.85, [
    [("Plus  ", {"bold": True, "color": BLUE}), ("AI package builder from free-text interests and booking history · priced budget fixes when a change goes over budget · recommended upgrades that still fit · undo and discard", {})],
], size=14, color=INK, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 2)
notes(s, "Four pillars: a curated start from the real PS-04 catalogue; live reshaping where the total is recomputed from the whole plan; the guide as a component with the mandatory availability check; and language plus channel choice. On top: an AI builder that only picks real catalogue packages, and budget fixes when the plan goes over budget.")

# ---------------------------------------------------------------- 3 · Architecture / technical approach
s = slide(LIGHT)
header(s, "Architecture · technical approach", "One TypeScript service, two channels, one trip engine")
layers = [
    ("Channels", ["Web app", "React 19 · Vite · Tailwind", "Telegram bot", "@wayypoint_Bot"]),
    ("API", ["tRPC + Zod", "Node 22 · Express", "Typed end to end,", "inputs validated"]),
    ("Trip engine", ["Pricing in integer paise", "Swaps · budget fixes", "Guide check", "Booking transaction"]),
    ("Data", ["PS-04 dataset (read-only)", "Canonical trips, itineraries,", "items, bookings", "+ app_* additions"]),
]
bw, gap = 2.65, 0.55
for i, (title, rows) in enumerate(layers):
    x = M + i * (bw + gap)
    card(s, x, 1.75, bw, 2.25, fill=PAPER if i != 2 else "EAF2FD", line=LINE if i != 2 else "9CC3F0")
    text(s, x + 0.2, 1.92, bw - 0.4, 0.4, title, size=18, bold=True, font=HEAD, color=BLUE if i == 2 else INK)
    text(s, x + 0.2, 2.42, bw - 0.4, 1.5, rows, size=13, color=SOFT, spacing=1.2)
    if i < 3:
        arrow(s, x + bw + 0.06, 2.72)
side = [
    ("AI — grounded, validated", "Sarvam sarvam-105b-conversations (OpenRouter fallback). Picks only real catalogue IDs; budget checked before and after the model; rule-based fallbacks."),
    ("Integrations", "SerpAPI Google Flights (live fares) · Sarvam mayura translation (cached) · Wikipedia photos · Railway deploy with a /data volume."),
]
for i, (title, body) in enumerate(side):
    x = M + i * ((W - 2 * M) / 2 + 0.12)
    card(s, x, 4.25, (W - 2 * M) / 2 - 0.12, 1.3)
    text(s, x + 0.25, 4.38, (W - 2 * M) / 2 - 0.6, 0.35, title, size=16, bold=True, font=HEAD, color=INK)
    text(s, x + 0.25, 4.78, (W - 2 * M) / 2 - 0.6, 0.75, body, size=12, color=SOFT, spacing=1.15)
decisions = ["One engine, two channels: a rule is enforced once", "Recompute, never accumulate: totals can't drift", "BEGIN IMMEDIATE + idempotency key: no double booking"]
dw = (W - 2 * M - 2 * 0.2) / 3
for i, d in enumerate(decisions):
    x = M + i * (dw + 0.2)
    card(s, x, 5.8, dw, 0.8, fill=NAVY, line=None)
    text(s, x + 0.2, 5.8, dw - 0.4, 0.8, d, size=13, bold=True, color=LIGHT, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
footer(s, 3)
notes(s, "The web app and the Telegram bot both call the same trip engine, so pricing, the guide check and booking behave identically everywhere. The engine reads the organisers' PS-04 dataset read-only and writes the shared model's canonical tables. Money is integer paise; the total is rebuilt from the plan on every change; booking is one SQLite transaction with a write lock and an idempotency key. The AI never sets a price.")

# ---------------------------------------------------------------- 4 · Expected MVP
s = slide(LIGHT)
header(s, "Expected MVP — 24-hour hackathon", "Every MVP item, built and wired end to end")
rows = [
    ("Package listing + ONE fully customisable package", "45 curated packages by theme; any one opens fully customisable", "MVP"),
    ("Swap hotel tier / activity / transfer / guide → itinerary and total update live", "Same day and slot kept; total rebuilt from the whole plan in paise", "MVP"),
    ("Guide as a component: language + specialisation + price", "Guide planner grid; day rate × date multiplier added to the total", "MVP"),
    ("Language preference changes packages, guides or content", "Profile sets app + guide language; guides filtered; 4 app languages", "MVP"),
    ("AI package-builder from free-text interests", "Sarvam; interests + booking history; real catalogue IDs only", "Stretch"),
    ("Save / share; full multilingual content", "Drafts, share links, PDF quotation, Telegram; 4-language copy", "Stretch"),
]
tbl = s.shapes.add_table(len(rows) + 1, 3, Inches(M), Inches(1.7), Inches(W - 2 * M), Inches(4.6)).table
widths = [5.0, 5.53, 1.6]
for c, wv in enumerate(widths):
    tbl.columns[c].width = Inches(wv)
tbl.rows[0].height = Inches(0.5)
for r in range(1, len(rows) + 1):
    tbl.rows[r].height = Inches(0.66)
for c, label in enumerate(["PS-04 expects", "What PackagePro does", "Status"]):
    cell = tbl.cell(0, c)
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb(NAVY)
    cell.text = label
    run = cell.text_frame.paragraphs[0].runs[0]
    run.font.size, run.font.bold, run.font.name = Pt(14), True, BODY
    run.font.color.rgb = rgb(LIGHT)
for r, (need, how, kind) in enumerate(rows, start=1):
    for c, value in enumerate([need, how, f"Built · {kind}"]):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(PAPER if r % 2 else "EEF3FA")
        cell.text = value
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.size, run.font.name = Pt(13), BODY
        run.font.bold = c != 1
        run.font.color.rgb = rgb(GREEN if c == 2 else INK if c == 0 else SOFT)
text(s, M, 6.5, W - 2 * M, 0.35, "Boundary rules enforced in the backend with tests: group size (min–max), BCP-47 languages, INR only, guide capacity, idempotent booking.", size=12, color=SOFT)
footer(s, 4)
notes(s, "Each MVP line from the statement, and where it lives in the product. Everything is wired end to end in code — the listing, the fully customisable package, live repricing on every swap, the guide as a priced component, and the language preference. Both stretch goals are built too: the AI package builder and save/share with full multilingual content.")

# ---------------------------------------------------------------- 5 · Mandatory enhancement — Guide Availability Check
s = slide(NAVY)
header(s, "Mandatory enhancement", "Guide Availability Check — never a static dropdown", dark=True)
steps = [
    ("Add a guide", "to a package with real trip dates"),
    ("Check every date", "is_available = 1 and slots_available − confirmed bookings > 0"),
    ("Refuse on a clash", "the clashing date is named"),
    ("Offer the nearest substitute", "same language + specialisation, free on every date, nearest (≤ 400 km), then smallest price change"),
    ("Reprice live", "price delta applied to the total; re-checked inside the booking transaction"),
]
for i, (title, body) in enumerate(steps):
    y = 1.75 + i * 0.95
    circ = card(s, M, y, 0.55, 0.55, fill=ORANGE, line=None, shape=MSO_SHAPE.OVAL)
    text(s, M, y, 0.55, 0.55, str(i + 1), size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, M + 0.8, y - 0.02, 5.8, 0.35, title, size=17, bold=True, font=HEAD, color=LIGHT)
    text(s, M + 0.8, y + 0.33, 5.8, 0.55, body, size=12, color=MUTED_ON_DARK, spacing=1.1)
# what the judges see
X = 7.3
text(s, X, 1.72, 5.4, 0.3, "WHAT THE JUDGES SEE", size=12, bold=True, color=ORANGE)
card(s, X, 2.1, 5.43, 1.45, fill=RED_BG, line=RED_LINE)
pill(s, X + 0.25, 2.28, 1.3, "REFUSED", "FDE3DE", RED)
text(s, X + 0.25, 2.7, 5.0, 0.4, "Meera Novak · Tamil · heritage", size=17, bold=True, font=HEAD, color=INK)
text(s, X + 0.25, 3.1, 5.0, 0.35, "Unavailable on 28 Sep 2026 — the date is named", size=13, color=RED)
arrow(s, X + 2.5, 3.65, w=0.4, h=0.45, color=ORANGE, shape=MSO_SHAPE.DOWN_ARROW)
card(s, X, 4.2, 5.43, 1.75, fill=GREEN_BG, line=GREEN_LINE)
pill(s, X + 0.25, 4.38, 2.1, "SUBSTITUTE OFFERED", "D3F1E3", GREEN)
text(s, X + 0.25, 4.8, 5.0, 0.4, "Arjun Nair · Tamil · heritage", size=17, bold=True, font=HEAD, color=INK)
text(s, X + 0.25, 5.2, 5.0, 0.7, ["Same language and specialisation · free on all 3 days", "Price delta applied · total repriced live"], size=13, color=GREEN, spacing=1.15)
text(s, X, 6.15, 5.43, 0.5, "Proof: tests/hardProof.guideAvailability.test.ts — same rule on web and Telegram, and day-by-day guide plans.", size=11, color=MUTED_ON_DARK)
footer(s, 5, dark=True)
notes(s, "This is the mandatory enhancement. Adding a guide checks every trip date against guide_availability and our own confirmed bookings. Meera speaks Tamil and specialises in heritage, but is unavailable on 28 September, so she is refused and the date is named. Arjun Nair — same language, same specialisation, free on all three days — is offered, and the total reprices live. The rule is enforced in the backend and proven by the hard-proof test; it is re-checked inside the booking transaction so a guide can't be double-booked.")

# ---------------------------------------------------------------- 6 · Demo
s = slide(LIGHT)
header(s, "Demo", "From search to booked trip in six steps")
demo = [
    ("Try the live demo", "Delhi → Thanjavur, 4 travellers, Tamil guide, ₹1,50,000 — live estimate: low · typical · high."),
    ("Pick a live flight", "Google Flights fares via SerpAPI; the package loads with its itinerary."),
    ("Customise", "Swap the hotel or an activity, add an extra, change the days — the total updates on every change."),
    ("Add a guide", "Meera Novak is refused on 28 Sep; Arjun Nair is offered with the new total."),
    ("Stay in budget", "Over budget? Priced fixes, or 'fit my budget automatically' in one tap."),
    ("Confirm & share", "Booking reference, PDF quotation in 4 languages — and the same flow in Telegram."),
]
cw, ch = (W - 2 * M - 2 * 0.3) / 3, 2.05
for i, (title, body) in enumerate(demo):
    x = M + (i % 3) * (cw + 0.3)
    y = 1.75 + (i // 3) * (ch + 0.3)
    card(s, x, y, cw, ch)
    card(s, x + 0.25, y + 0.25, 0.5, 0.5, fill=BLUE, line=None, shape=MSO_SHAPE.OVAL)
    text(s, x + 0.25, y + 0.25, 0.5, 0.5, str(i + 1), size=15, bold=True, color=LIGHT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.95, y + 0.3, cw - 1.2, 0.45, title, size=17, bold=True, font=HEAD, color=INK)
    text(s, x + 0.25, y + 0.9, cw - 0.5, 1.1, body, size=13, color=SOFT, spacing=1.2)
text(s, M, 6.55, W - 2 * M, 0.35, [[("Try it  ", {"bold": True, "color": BLUE}), ("[live URL] · /how-it-works guided tour · Telegram @wayypoint_Bot → /demo", {})]], size=13, color=INK)
footer(s, 6)
notes(s, "Run 'pnpm db:reset' before the demo so the demo guide's slots are free. Click 'Try the live demo', pick the cheapest flight, swap the hotel, open the guide planner and pick Meera Novak — refused on 28 September, Arjun offered with the new total. Confirm to get a booking reference and download the PDF. Then show the same refusal in the Telegram bot.")

# ---------------------------------------------------------------- 7 · Proof
s = slide(BLUE)
header(s, "Proof", "Backed by code, tests and the validator", dark=True)
proof = [
    ("61", "automated tests", "including the hard-proof guide test, concurrency and idempotent booking"),
    ("PASS", "organisers' validator", "our canonical rows merged with PS-04 pass validate_conformance.py"),
    ("120", "bookable guides", "each with a 30-day calendar; plus 45 packages and 1,200 traveller profiles from PS-04"),
    ("4 · 2", "languages · channels", "English, Hindi, Tamil, Telugu on web and Telegram"),
]
cw = (W - 2 * M - 3 * 0.25) / 4
for i, (big, label, body) in enumerate(proof):
    x = M + i * (cw + 0.25)
    card(s, x, 1.9, cw, 3.6, fill="0A5AAD", line="3C86D6")
    text(s, x + 0.25, 2.15, cw - 0.5, 1.0, big, size=40, bold=True, font=HEAD, color=LIGHT, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.25, 3.25, cw - 0.5, 0.45, label, size=15, bold=True, color="FFE3CC")
    text(s, x + 0.25, 3.75, cw - 0.5, 1.6, body, size=13, color="E3EEFB", spacing=1.2)
text(s, M, 5.85, W - 2 * M, 0.6, "Boundary rules live in the backend: group size, BCP-47 languages, INR only, guide capacity. Money is integer paise; itinerary lines add up to the total to the paisa.", size=14, color=LIGHT)
footer(s, 7, dark=True)
notes(s, "pnpm verify runs the type check, every test suite and the production build before every push. pnpm conformance merges our canonical rows with the PS-04 dataset and runs the organisers' validator — it prints PASS. Everything shown is backed by code.")

# ---------------------------------------------------------------- 8 · Impact and benefits
s = slide(LIGHT)
header(s, "Impact & benefits", "Better for travellers, guides and tour operators")
impact = [
    ("Travellers", ["Build the trip they want, in their language", "See every line of the price", "Never booked with a guide who isn't free", "Stay in budget with one-tap fixes"]),
    ("Local guides", ["Bookings respect real availability and slots", "Matched on language and specialisation", "Nearby guides get offered as substitutes", "Paid per day at each date's rate"]),
    ("Tour operators", ["Fewer manual re-quotes and call-backs", "No double-booked guides", "Upsell with upgrades that still fit", "One engine for web and chat"]),
]
cw = (W - 2 * M - 2 * 0.3) / 3
for i, (title, points) in enumerate(impact):
    x = M + i * (cw + 0.3)
    card(s, x, 1.75, cw, 4.6, radius=0.03)
    card(s, x, 1.75, cw, 0.09, fill=[BLUE, ORANGE, GREEN][i], line=None, radius=0, shape=MSO_SHAPE.RECTANGLE)
    text(s, x + 0.3, 2.05, cw - 0.6, 0.5, title, size=22, bold=True, font=HEAD, color=INK)
    text(s, x + 0.3, 2.75, cw - 0.6, 3.4, [[("•  ", {"bold": True, "color": [BLUE, "C2410C", GREEN][i]}), (p, {})] for p in points], size=15, color=SOFT, spacing=1.3)
footer(s, 8)
notes(s, "Travellers get control and transparency in their own language. Guides are only booked when they are really free, and are matched on language and specialisation. Operators stop re-quoting by hand, never double-book a guide, and can upsell with upgrades that still fit the budget.")

# ---------------------------------------------------------------- 9 · Business model & scaling
s = slide(NAVY)
header(s, "Business model & scaling", "How PackagePro earns — and grows beyond one city", dark=True)
revenue = [
    ("Booking commission", "[__%] of each confirmed package"),
    ("Guide marketplace fee", "[__%] per guide-day booked"),
    ("White-label for agencies", "[₹__ / month] per agency, their brand and catalogue"),
    ("Upgrades & add-ons", "margin on recommended upgrades that fit the budget"),
]
text(s, M, 1.75, 5.8, 0.35, "REVENUE", size=12, bold=True, color=ORANGE)
for i, (title, body) in enumerate(revenue):
    y = 2.15 + i * 1.02
    card(s, M, y, 5.9, 0.88, fill=CARD_DARK, line=LINE_DARK)
    text(s, M + 0.25, y + 0.1, 5.4, 0.35, title, size=16, bold=True, font=HEAD, color=LIGHT)
    text(s, M + 0.25, y + 0.47, 5.4, 0.35, body, size=12, color=MUTED_ON_DARK)
X = 7.1
text(s, X, 1.75, 5.6, 0.35, "SCALING PATH", size=12, bold=True, color=ORANGE)
scale = [
    ("Today", "One Railway service, SQLite, canonical tables, Telegram polling"),
    ("Next", "Postgres with the same schema; stateless engine on several instances"),
    ("Then", "Telegram webhook + WhatsApp; UPI payments and GST on the quotation"),
    ("Later", "More cities, guides and partner inventory through the same engine"),
]
for i, (stage, body) in enumerate(scale):
    y = 2.15 + i * 1.02
    card(s, X, y, 1.2, 0.88, fill=ORANGE, line=None)
    text(s, X, y, 1.2, 0.88, stage, size=14, bold=True, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    card(s, X + 1.3, y, W - M - X - 1.3, 0.88, fill=CARD_DARK, line=LINE_DARK)
    text(s, X + 1.5, y, W - M - X - 1.7, 0.88, body, size=13, color=LIGHT, anchor=MSO_ANCHOR.MIDDLE)
text(s, M, 6.35, W - 2 * M, 0.45, "Thank you — questions?", size=22, bold=True, font=HEAD, color=LIGHT, align=PP_ALIGN.CENTER)
footer(s, 9, dark=True)
notes(s, "Revenue: a commission on each booked package, a marketplace fee on guide-days, a white-label licence for agencies, and margin on upgrades. Fill in the rates before presenting. Scaling: today one service with SQLite by design; next Postgres with the same canonical schema and a stateless engine; then a Telegram webhook, WhatsApp, UPI payments and GST; then more cities and partner inventory — the engine code does not change.")

prs.save(OUT)
print(f"saved {OUT} ({len(prs.slides)} slides)")
